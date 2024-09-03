import argparse
import csv
from collections import namedtuple
from typing import List

from openpyxl.workbook import Workbook
from openpyxl.worksheet.table import Table

from parser import ParserFactory, Row
from utils import file_exists

HEADERS = [
    "date",
    "descrição",
    "valor",
    "conta",
    "category",
    "installment",
]  # Mobills has an inconsistent localization config, requiring a mix of English and Portuguese in headers


def read_sheet(sheet, template) -> List[Row]:
    parser_class = ParserFactory.get_parser(template)
    parser = parser_class()

    rows = []
    for row in sheet:
        try:
            data = parser.parse(row)
            rows.append(data)

        except ValueError:
            continue

    return rows


def _add_headers(ws):
    ws.append(HEADERS)


def _add_content(ws, rows: List[Row]):
    RowData = namedtuple(
        "RowNamedTuple",
        ["date", "description", "amount", "account", "category", "installment"],
    )
    for row in rows:
        ws.append(
            RowData(
                date=row.date,
                description=f"{row.description}",
                amount=row.amount,
                account="",
                category="",
                installment=row.installment,
            )
        )


def write_sheet(filename, rows: List[Row]):
    wb = Workbook()
    ws = wb.active

    _add_headers(ws)
    _add_content(ws, rows)

    # Define the range for the table
    table_range = f"A1:F{ws.max_row}"

    # Create a table
    tab = Table(displayName="TransactionTable", ref=table_range)

    # Add the table to the worksheet
    ws.add_table(tab)

    # Save the workbook
    name, _ = filename.split(".")
    wb.save(f"{name}.xlsx")
    print(f"Saved '{name}.xlsx'")


def run(filename: str, template: str):
    file_exists(filename)  # Validate if the file exists

    with open(filename, newline="", encoding="utf-8") as csv_file:
        csv_reader = csv.reader(csv_file)
        content = read_sheet(csv_reader, template)
        write_sheet(filename, content)


if __name__ == "__main__":
    arg_parser = argparse.ArgumentParser(
        description="Process a CSV file using a specified template."
    )

    arg_parser.add_argument("filename", help="The name of the file to process")
    arg_parser.add_argument(
        "template", help="The bank template, options: inter, santander"
    )

    args = arg_parser.parse_args()

    run(args.filename, args.template)
